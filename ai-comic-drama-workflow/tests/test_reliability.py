from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

from ai_comic_drama_workflow.capabilities import probe_capabilities
from ai_comic_drama_workflow.exporter import build_delivery_archive
from ai_comic_drama_workflow.ingest import classify_source, extract_text
from ai_comic_drama_workflow.kernel import WorkflowKernel
from ai_comic_drama_workflow.layout import P09, P10, SHOT_TIMELINE_INDEX
from ai_comic_drama_workflow.pipeline import canonical_prompt_for_shot
from ai_comic_drama_workflow.storage import ProjectStore
from ai_comic_drama_workflow.utils import SCHEMA_VERSION
from tests.helpers import artifact_for_stage, drive_to_completion
from tests.test_v2_workflow import advance_until


class ReliabilityTests(unittest.TestCase):
    def test_invalid_final_report_blocks_automatic_and_explicit_export(self):
        with tempfile.TemporaryDirectory() as temp:
            kernel = WorkflowKernel.initialize(Path(temp) / "project", ["小说正文"])
            with patch("ai_comic_drama_workflow.pipeline.validate_project", return_value={"valid": False, "errors": [{"code": "injected"}]}):
                with self.assertRaises(AssertionError):
                    drive_to_completion(kernel, storyboard_durations=[4])
            self.assertNotEqual("complete", kernel.state["status"])
            self.assertFalse(list(kernel.store.root.rglob("*.zip")))
            # An already completed project must also revalidate on explicit export.
            kernel.state  # no cached completion result is used
            kernel.run()
            drive_to_completion(kernel, storyboard_durations=[4])
            prompt = next(kernel.store.path(P10).rglob("*.txt"))
            prompt.write_text("tampered", encoding="utf-8")
            with self.assertRaises(ValueError):
                build_delivery_archive(kernel.store)

    def test_revision_runs_back_to_delivery(self):
        with tempfile.TemporaryDirectory() as temp:
            kernel = WorkflowKernel.initialize(Path(temp) / "project", ["小说正文"])
            drive_to_completion(kernel, storyboard_durations=[4, 4])
            original = kernel.store.read(SHOT_TIMELINE_INDEX)
            kernel.invalidate_from("spatial_blocking")
            self.assertEqual("complete", drive_to_completion(kernel, storyboard_durations=[4, 4])["status"])
            self.assertGreater(kernel.store.read(SHOT_TIMELINE_INDEX)["revision"], original["revision"])

    def test_second_invalid_shot_leaves_no_partial_batch(self):
        with tempfile.TemporaryDirectory() as temp:
            kernel = WorkflowKernel.initialize(Path(temp) / "project", ["小说正文"])
            task = advance_until(kernel, "shot_timeline_specs", durations=[4, 4])["task_envelope"]
            artifact = artifact_for_stage(kernel, task, storyboard_durations=[4, 4])
            artifact["shots"][1]["scene_id"] = "DOES-NOT-EXIST"
            result = {"schema_version": SCHEMA_VERSION, "result_id": "bad-batch", "project_id": task["project_id"],
                      "stage_id": task["stage_id"], "context_fingerprint": task["context_fingerprint"],
                      "status": "draft", "artifact": artifact, "findings": [], "provenance": []}
            with self.assertRaises(RuntimeError):
                kernel.submit(result)
            self.assertEqual([], list(kernel.store.path(f"{P09}/shots").glob("*.json")))
            self.assertFalse(kernel.store.exists(SHOT_TIMELINE_INDEX))

    def test_performance_change_changes_execution_prompt(self):
        shot = {"character_tracks": [{"character_id": "C", "trajectory": [{"time_s": 0, "gaze": "door"}],
                                       "emotion_events": [{"visible_cues": ["frown"], "trigger": "knock"}]}]}
        modified = deepcopy(shot)
        modified["character_tracks"][0]["trajectory"][0]["gaze"] = "partner"
        modified["character_tracks"][0]["emotion_events"][0]["visible_cues"] = ["smile"]
        self.assertNotEqual(canonical_prompt_for_shot(shot)[0], canonical_prompt_for_shot(modified)[0])

    def test_declarations_are_not_executable_capability(self):
        report = probe_capabilities(environment={"AI_COMIC_DRAMA_TRUE_3D_PROVIDER": "invented",
            "AI_COMIC_DRAMA_TRUE_3D_STATUS": "available", "AI_COMIC_DRAMA_TRUE_3D_EXECUTABLE": "true"}, which=lambda _: None)
        self.assertNotEqual("available", report["true_3d"]["status"])
        self.assertFalse(report["true_3d"]["executable"])

    def test_versioned_shot_is_not_project_migration(self):
        payload = {"schema_version": "2.0", "project_id": "P", "artifact_type": "shot-timeline-spec"}
        self.assertEqual("storyboard", classify_source(Path("input.json"), json.dumps(payload)))

    def test_docx_keeps_table_in_document_order(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "shots.docx"
            document = Document()
            document.add_paragraph("before")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "SHOT 01"
            document.add_paragraph("after")
            document.save(path)
            text, status = extract_text(path)
            self.assertEqual("extracted", status)
            self.assertLess(text.index("before"), text.index("SHOT 01"))
            self.assertLess(text.index("SHOT 01"), text.index("after"))

    def test_crash_write_set_is_recovered(self):
        with tempfile.TemporaryDirectory() as temp:
            store = ProjectStore(temp)
            store.write_text("existing.txt", "before")
            script = "from ai_comic_drama_workflow.storage import ProjectStore; import os,sys; s=ProjectStore(sys.argv[1]); c=s.transaction(); c.__enter__(); s.write_text('existing.txt','after'); s.write_text('new.txt','partial'); os._exit(7)"
            result = subprocess.run([sys.executable, "-c", script, temp], env={**__import__('os').environ,
                "PYTHONPATH": str(Path(__file__).resolve().parents[1] / "src")}, capture_output=True)
            self.assertEqual(7, result.returncode)
            with store.transaction():
                self.assertEqual("before", store.path("existing.txt").read_text())
                self.assertFalse(store.exists("new.txt"))


if __name__ == "__main__":
    unittest.main()
